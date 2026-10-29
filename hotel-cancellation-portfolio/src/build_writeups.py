"""Build provisional portfolio text from recorded results, never invented metrics."""

from pathlib import Path
import json
import pandas as pd


def build(root):
    root = Path(root)
    r = json.loads((root / "reports/results.json").read_text())
    a = json.loads((root / "reports/audit.json").read_text())
    t = pd.read_csv(root / "reports/model_comparison.csv").set_index("experiment")
    v = json.loads((root / "private/submission_validation.json").read_text())
    best = r["local_mean_auc"]
    std = r["local_std_auc"]
    grouped = r["grouped_diagnostic_mean_auc"]

    def score(n):
        return float(t.loc[n, "mean_auc"])

    table = "| Experiment | Mean local AUC | Fold SD | Runtime (seconds) |\n|---|---:|---:|---:|\n"
    for n, row in t.iterrows():
        table += (
            f"| {n} | {row.mean_auc:.6f} | {row.std_auc:.6f} | {row.seconds:.1f} |\n"
        )
    summary = f"""# What changed the results

All main experiments used identical five-fold stratified validation (shuffle=True, random_state=835). These are local ROC-AUC results, not Kaggle results or percent accuracy.

{table}

- Treating agent/company as categories instead of numeric amounts changed logistic mean AUC by {score('02_lr_categorical_ids')-score('01_provisional_lr_numeric_ids'):+.6f}.
- Adding four simple row features changed logistic mean AUC by {score('03_lr_features')-score('02_lr_categorical_ids'):+.6f}. This is a small difference relative to fold variation; it is not proof of a reliable improvement in future data.
- Changing logistic C from 1 to 0.1 changed mean AUC by {score('04_lr_features_C01')-score('03_lr_features'):+.6f}. A smaller C means stronger regularization.
- Increasing boosting leaves from 15 to 31 changed mean AUC by {score('07_hgb_leaf31')-score('06_hgb_leaf15'):+.6f}.
- Dropping the four timing-uncertain fields from the 15-leaf boosting reference changed mean AUC by {score('08_hgb_conservative')-score('06_hgb_leaf15'):+.6f}.
- Removing country from the same boosting reference changed mean AUC by {score('09_hgb_no_country')-score('06_hgb_leaf15'):+.6f}.
- Removing deposit type from the same reference changed mean AUC by {score('10_hgb_no_deposit')-score('06_hgb_leaf15'):+.6f}.

Selected candidate: **{r['selected_model']}**, mean AUC **{best:.6f}**, fold SD **{std:.6f}**. Selection used the first seven predefined candidate experiments; the three feature ablations are diagnostics, not extra winner candidates. Settings: `{json.dumps(r['settings'],sort_keys=True)}`.

The selected model's duplicate-profile grouped diagnostic scored **{grouped:.6f}** mean AUC, compared with **{best:.6f}** on random stratified folds. This difference reflects a different validation design, not a paired isolated model comparison. Grouping exact profiles still cannot identify the same guest or hotel booking across imperfectly matching records.

Pooled out-of-fold AUC is {r['pooled_oof_auc']:.6f}; it differs from the average of fold AUCs because it pools scores across separately trained models. Brier score is {r['brier_score']:.6f}; smaller is better. The ROC and calibration charts use held-out predictions. Descriptive country/deposit charts use labeled training rows and do not claim causality.

Model selection reused validation folds and introduces optimism; no nested CV or untouched holdout score is claimed. The official starter, prediction-time definitions and course-method permissions remain unverified. No public or private Kaggle AUC is available.
"""
    (root / "reports/interpretation.md").write_text(summary)
    readme = f"""# Team HH — Hotel Booking Cancellation

An ISOM 835 competition project at Suffolk University, prepared as a beginner-friendly portfolio draft. Predict a cancellation probability for each booking using only competition-provided data.

**Status:** runnable provisional workflow with executed local results. Official starter, professor instructions, course notebooks, project guide, sample submission and memo rubric are still needed. This is not a claim that every assignment requirement is complete.

## Verified local results

Selected model: **{r['selected_model']}**. Five-fold stratified local ROC-AUC: **{best:.6f} ± {std:.6f}** (fold standard deviation, not a confidence interval). Grouped duplicate-profile diagnostic mean AUC: **{grouped:.6f}**. Public Kaggle AUC: not submitted. Private Kaggle AUC: unavailable. ROC-AUC describes ranking and must not be called percent accuracy.

{table}

## Data and methods

Training: {a['train_shape'][0]:,} bookings, including is_canceled. Test: {a['test_shape'][0]:,} bookings with hidden outcomes. Cancellation rate in training: {a['cancellation_rate']:.2%}. {a['train_duplicate_predictors_excluding_id']:,} training rows repeat an earlier predictor profile after excluding ID. Data and row-level outputs are deliberately not included in this repository.

Compare provisional logistic regression, improved categorical-ID logistic regression, random forest and histogram gradient boosting. Use identical stratified folds (5 splits, shuffle=True, random_state=835), fold-fitted preprocessing, train-only rare-category handling and four simple row-level features. A small predefined search tests two logistic regularization strengths and two boosting leaf counts. Country, deposit type and timing-uncertain fields get separate diagnostic ablations. Winner chosen using local CV, never public leaderboard tuning.

## Run in Google Colab

1. Open notebooks/TeamHH_Hotel_Cancellation.ipynb in Colab.
2. Select Runtime → Run all. Upload your own train.csv and test.csv, plus sample_submission.csv when available.
3. Read the audit, model comparison and charts. The notebook fits 55 fold models and one final model; runtime depends on the machine.
4. The checked TeamHH_submission.csv is written to private/. Change DOWNLOAD_SUBMISSION to True to download it in Colab. Never commit predictions or data.

The notebook embeds the analysis code, so it does not need a GitHub checkout to run. It has numbered code cells and short explanations. It was tested locally with the recorded library versions; an actual hosted Colab run is a separate check.

## Run locally

Python 3.12 was used for recorded results. From the project folder:

```bash
python -m pip install -r requirements.txt
# Place competition CSVs in data/ first.
OPENBLAS_NUM_THREADS=2 OMP_NUM_THREADS=2 python src/workflow.py --data-dir data --root . --team TeamHH
python src/build_writeups.py
```

## What to read

- docs/requirements_checklist.md: completed work and exact missing materials.
- docs/leakage_audit.md: prediction-time uncertainty, duplicate profiles and leakage controls.
- reports/interpretation.md: what helped and what did not, with actual score differences.
- reports/fold_scores.csv: every main fold score, settings, runtime, reason and warnings.
- reports/grouped_sensitivity.csv: separate grouped-fold diagnostic.
- reports/results.json: model settings, data hashes, package version and unverified items.
- docs/memo_draft.pdf and memo_draft.md: one-page memo draft; rubric not yet supplied.
- docs/presentation_outline.md and learning_notes.md: three-minute talk and plain-English explanations.
- docs/linkedin_draft.md: review-only post with verified results and credit placeholders.
- docs/submission_instructions.md: deadlines, submission limits and file identity checks.
- docs/github_handoff.md: private-repository upload instructions and current access blocker.

## Limitations

The main candidate assumes provided features describe a usable pre-outcome snapshot. Their initial-booking availability is not confirmed. History variables need strict as-of timestamps. The conservative ablation cannot establish deployment readiness. Country/deposit associations are not causal and raise interpretation/fairness concerns. Random folds may be optimistic due to repeated profiles; grouped sensitivity is reported separately. Selecting on the same CV folds adds selection optimism. No future-time, real hotel deployment, business savings or guaranteed competition score is claimed.

## Team contributions and AI disclosure

Team: **Team HH**. Individual names, roles and contributions must be supplied by the team before sharing. Do not represent the prepared code or analysis as independently authored. AI assisted with data auditing, Python code, experiment design, interpretation and writing drafts. The student/team must review the code, understand the methods, confirm course permissions and record their own changes, decisions and final submission responsibilities. Only claim personal work actually performed.

## Sharing and submission

This project is prepared for review. Portfolio destination: [https://github.com/Hemanth-55/Hemanth_Test/tree/team-hh-hotel-cancellation/hotel-cancellation-portfolio](https://github.com/Hemanth-55/Hemanth_Test/tree/team-hh-hotel-cancellation/hotel-cancellation-portfolio), a dedicated folder on the team-hh-hotel-cancellation branch of the existing public Hemanth-55/Hemanth_Test repository. The GitHub integration denied new-repository creation, so the existing connected repository is used. Competition data, per-booking predictions and credentials are excluded. Confirm course sharing rules before redistributing this work. No LinkedIn post, Kaggle upload, Canvas upload or deployment has been made. The competition data, predictions, OOF arrays, credentials and original uploaded materials remain excluded. No data license or permission to redistribute is asserted. Do not apply a license to the competition data.
"""
    (root / "README.md").write_text(readme)
    memo = f"""Team HH — ISOM 835 Hotel Cancellation | MEMO DRAFT
To: ISOM 835 instructor | From: Team HH (individual names pending)
Subject: Local validation and provisional competition submission

Purpose. We estimated cancellation probabilities for 35,817 test bookings using 83,573 labeled bookings. The training cancellation rate was {a['cancellation_rate']:.2%}. ROC-AUC measures ranking; it is not percent accuracy.

Method. We compared a provisional logistic reference, improved logistic regression, random forest and histogram gradient boosting on the same five stratified folds (shuffle=True, random_state=835). Imputation, scaling and category handling were fitted within training folds. Agent/company IDs were tested as categories. Four simple booking features and a small predefined parameter search were evaluated. No external data or hidden-label recovery was used.

Findings. The selected model, {r['selected_model']}, achieved local mean AUC {best:.6f} (fold SD {std:.6f}). Categorical IDs changed logistic AUC from {score('01_provisional_lr_numeric_ids'):.6f} to {score('02_lr_categorical_ids'):.6f}. Added features produced {score('03_lr_features'):.6f}; stronger regularization produced {score('04_lr_features_C01'):.6f}. Full scores and settings are in reports/. Public and private Kaggle scores are unavailable; no submission was made.

Risks. There were {a['train_duplicate_predictors_excluding_id']:,} repeated training predictor profiles. The selected model's grouped-profile diagnostic mean AUC was {grouped:.6f}; it uses different folds. Some predictors may be recorded after initial booking, so a conservative timing ablation was included. Country and deposit patterns are associations, not causes. Selection on CV folds adds optimism; no deployment or savings are claimed.

Submission. TeamHH_submission.csv passed row, ID, column and probability checks and has a recorded SHA256. The actual sample template is missing. First Kaggle submission is required by November 2; final predictions by November 9, 2026, 4:30 PM (confirm professor timezone). At most five Kaggle uploads per day; Canvas must receive the exact last Kaggle file.

Disclosure and outstanding review. AI assisted code, auditing, experiments and writing. Team members must verify understanding, permissions and individual contributions. The official starter, guide, professor/course materials and separate memo rubric are missing. This is not the official baseline or a rubric-complete memo. Review these documents before final submission.
"""
    (root / "docs/memo_draft.md").write_text(memo)
    # Render an actual single-page PDF. Keep text wrapped and below the page boundary.
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import letter
    import textwrap

    c = canvas.Canvas(str(root / "docs/memo_draft.pdf"), pagesize=letter)
    c.setTitle("Team HH - ISOM 835 memo draft - rubric pending")
    y = 750
    for p in memo.split("\n"):
        for line in textwrap.wrap(p, width=106) or [""]:
            if y < 44:
                raise ValueError("Memo overflow: shorten before publishing")
            c.setFont("Helvetica", 9.3)
            c.drawString(42, y, line)
            y -= 12
    c.showPage()
    c.save()
    # Word copy for the student to edit after the rubric arrives.
    from docx import Document
    from docx.shared import Inches, Pt

    d = Document()
    section = d.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.6)
    style = d.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(9)
    style.paragraph_format.space_after = Pt(4)
    for p in memo.split("\n"):
        if p:
            d.add_paragraph(p)
    d.save(root / "docs/memo_draft.docx")
    presentation = f"""# Three-minute presentation outline — Team HH

**0:00–0:25 — Problem.** “We predict the chance that a hotel booking will be canceled. We have 83,573 labeled bookings and 35,817 hidden-outcome bookings. We evaluate ranking with ROC-AUC, not percent accuracy.”

**0:25–0:55 — Audit and leakage.** Show reports/audit.json and mention the {a['cancellation_rate']:.2%} cancellation rate. “IDs are excluded. Duplicate profiles can make random validation optimistic. Some booking fields may change after booking, so their timing needs confirmation.”

**0:55–1:25 — Fair comparison.** Show reports/model_comparison.png. Explain five folds in one sentence: “Each booking is evaluated by a model that did not train on that row.” Mention fold-fitted preprocessing, categorical agent/company IDs and the limited predefined search.

**1:25–2:00 — Actual findings.** “Categorical IDs changed logistic AUC from {score('01_provisional_lr_numeric_ids'):.4f} to {score('02_lr_categorical_ids'):.4f}. Our selected {r['selected_model']} scored {best:.4f} mean local AUC with fold SD {std:.4f}. Grouped-profile validation gave {grouped:.4f}; this is a different test of generalization.” Read reports/interpretation.md for features and settings that did not help.

**2:00–2:30 — Meaning and limits.** Show validation_roc.png or validation_calibration.png. “Good ranking does not guarantee perfectly calibrated probabilities. Deposit/country patterns are associations, not causal claims. There are no Kaggle or deployment results yet.”

**2:30–3:00 — Handoff and credit.** “We checked the 35,817 probabilities, IDs, columns and file hash. Next we must review the official starter, course methods, sample and memo rubric, then submit the exact same final file to Kaggle and Canvas. AI assisted the draft; team members will explain their verified contributions.” Add confirmed teammate names and roles before speaking.
"""
    (root / "docs/presentation_outline.md").write_text(presentation)
    linkedin = f"""# LinkedIn draft — review before posting

I’m building a hotel booking cancellation project for ISOM 835 at Suffolk University with Team HH.

The task is to predict cancellation probabilities using competition-provided booking data. The project compares logistic regression, random forest and histogram gradient boosting using the same five cross-validation folds, with preprocessing learned only from each training fold.

In the executed local analysis, the selected model achieved mean ROC-AUC {best:.4f} (fold standard deviation {std:.4f}). A separate diagnostic that kept identical booking profiles together achieved mean AUC {grouped:.4f}. These are local validation results; no Kaggle score or deployment result is claimed.

One useful finding: agent and company codes behave like category IDs, not numeric amounts. Another lesson is that repeated profiles and the timing of booking fields can change how we interpret validation scores. Country and deposit patterns are associations, not causal explanations.

AI assisted with code, analysis and writing drafts. I’m reviewing the methods and course requirements before final submission and public sharing.

Team credit: Team HH — [ADD CONFIRMED NAMES AND EACH PERSON'S ACTUAL CONTRIBUTION]. My contribution: [DESCRIBE ONLY WORK I PERSONALLY COMPLETED AND CAN EXPLAIN].

Repository: [ADD LINK AFTER PRIVATE UPLOAD, SHARING-RULE REVIEW AND PUBLICATION DECISION]. Competition data and row-level predictions will not be shared.

#Python #MachineLearning #DataAnalytics #SuffolkUniversity

Review note (remove before posting): Professor instructions, official starter, course notebooks, project guide, sample submission and memo rubric are still missing. Do not imply this draft completes all graded requirements. Do not post until teammates approve credit and sharing is permitted.
"""
    (root / "docs/linkedin_draft.md").write_text(linkedin)
    (root / "docs/final_file_identity.md").write_text(
        f"""# Candidate file identity

File: `{v['file']}` (private download, excluded from GitHub).\nRows: {v['rows']:,}. Columns: `id,is_canceled`.\nSHA256: `{v['sha256']}`.\nSample template verified: {v['sample_submission_verified']}.\nThis is a provisional locally generated candidate, not a record of Kaggle/Canvas submission.\n"""
    )


if __name__ == "__main__":
    build(Path(__file__).resolve().parents[1])
